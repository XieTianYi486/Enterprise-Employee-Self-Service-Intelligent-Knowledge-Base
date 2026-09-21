# ============================================================
# 文档 Ingestion 流水线
# 编排：解析 → 分块 → 向量化 → ChromaDB + BM25 双索引
# 与旧项目通过 LangChain Loader 的方式区分开
# ============================================================

import os
import uuid
import hashlib
from pathlib import Path
from typing import List, Dict, Tuple, Optional

from app.core.config import settings
from app.core.exceptions import DocumentProcessException
from app.rag.chunker.semantic_chunker import get_chunker
from app.rag.embeddings.dashscope_embeddings import get_embeddings
from app.db.chroma_client import get_vector_store
from app.rag.retrievers.bm25_retriever import get_bm25_retriever, save_bm25_index


class DocumentParser:
    """
    文档解析器
    直接使用原生库（不使用 LangChain Loader）
    """

    @staticmethod
    def parse_pdf(file_path: str) -> Tuple[str, int]:
        """
        解析 PDF 文档
        使用 PyMuPDF (fitz) 直接提取文本
        """
        import fitz  # PyMuPDF

        text_parts = []
        page_count = 0

        with fitz.open(file_path) as doc:
            page_count = doc.page_count
            for page_num, page in enumerate(doc):
                page_text = page.get_text("text")
                if page_text.strip():
                    # 添加页码标记
                    text_parts.append(f"[第{page_num + 1}页]\n{page_text}")

        full_text = "\n\n".join(text_parts)
        return full_text, page_count

    @staticmethod
    def parse_docx(file_path: str) -> Tuple[str, int]:
        """
        解析 Word 文档
        使用 python-docx 提取文本和表格
        """
        from docx import Document

        doc = Document(file_path)
        text_parts = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

        # 逐段落提取
        for para in doc.paragraphs:
            if para.text.strip():
                # 识别标题样式
                if para.style.name.startswith("Heading"):
                    level = para.style.name.split()[-1]
                    prefix = "#" * min(int(level), 3)
                    text_parts.append(f"{prefix} {para.text}")
                else:
                    text_parts.append(para.text)

        # 提取表格
        for table in doc.tables:
            table_text = _table_to_markdown(table)
            if table_text:
                text_parts.append(table_text)

        full_text = "\n\n".join(text_parts)
        # 估算页数（约 1500 字符/页）
        estimated_pages = max(1, len(full_text) // 1500)
        return full_text, estimated_pages

    @staticmethod
    def parse_xlsx(file_path: str) -> Tuple[str, int]:
        """
        解析 Excel 文档
        使用 openpyxl 提取，每个 sheet 作为一节
        """
        import openpyxl

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        text_parts = []
        total_rows = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"## {sheet_name}")

            rows_data = []
            for row in ws.iter_rows(values_only=True):
                row_values = [
                    str(cell) if cell is not None else ""
                    for cell in row
                ]
                if any(v.strip() for v in row_values):
                    rows_data.append(row_values)
                    total_rows += 1

            if rows_data:
                text_parts.append(_rows_to_markdown_table(rows_data))

        wb.close()
        full_text = "\n\n".join(text_parts)
        estimated_pages = max(1, total_rows // 40)  # 约 40 行/页
        return full_text, estimated_pages

    @staticmethod
    def parse_markdown(file_path: str) -> Tuple[str, int]:
        """解析 Markdown 文档（保留原始格式）"""
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

        estimated_pages = max(1, len(text) // 1500)
        return text, estimated_pages

    @staticmethod
    def parse_txt(file_path: str) -> Tuple[str, int]:
        """解析纯文本（自动检测编码）"""
        # 尝试 UTF-8，失败则用 GBK
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="gbk") as f:
                text = f.read()

        estimated_pages = max(1, len(text) // 1500)
        return text, estimated_pages

    @classmethod
    def parse(cls, file_path: str, file_type: str) -> Tuple[str, int]:
        """
        根据文件类型自动选择解析器

        参数:
            file_path: 文件路径
            file_type: 文件类型 (pdf/docx/xlsx/md/txt)

        返回:
            (full_text, page_count)
        """
        parsers = {
            "pdf": cls.parse_pdf,
            "docx": cls.parse_docx,
            "xlsx": cls.parse_xlsx,
            "md": cls.parse_markdown,
            "txt": cls.parse_txt,
        }

        parser = parsers.get(file_type)
        if not parser:
            raise DocumentProcessException(f"不支持的文件类型: {file_type}")

        try:
            return parser(file_path)
        except Exception as e:
            raise DocumentProcessException(
                f"文档解析失败 ({file_type}): {str(e)}"
            )


# ==================== Ingestion 流水线 ====================


def process_document(file_path: str, file_type: str, document_id: int,
                     title: str, category_id: Optional[int] = None,
                     security_level: int = 1) -> Tuple[int, List[str], List[Dict]]:
    """
    完整的文档处理流水线

    流程：
    1. 解析文档 → 提取全文文本
    2. 语义分块 → chunk 列表
    3. 去重（SHA256）
    4. 向量化（百炼 Embedding）
    5. 写入 ChromaDB（向量索引）
    6. 写入 BM25（关键词索引）
    7. 持久化 BM25 索引到磁盘

    参数:
        file_path: 文档文件路径
        file_type: 文件类型
        document_id: 文档数据库 ID
        title: 文档标题
        category_id: 分类 ID
        security_level: 密级

    返回:
        (chunk_count, chunk_ids, chunk_records)
        chunk_records 包含 [{"content": str, "chapter": str, "page": int, "token_count": int}, ...]
    """
    chunker = get_chunker()
    embeddings = get_embeddings()
    vector_store = get_vector_store()
    bm25_retriever = get_bm25_retriever()

    # ===== 第1步：解析文档 =====
    full_text, page_count = DocumentParser.parse(file_path, file_type)

    if not full_text or not full_text.strip():
        raise DocumentProcessException("文档内容为空，无法处理")

    # ===== 第2步：语义分块 =====
    raw_chunks = chunker.chunk_text(
        text=full_text,
        title=title,
        page_start=1,
        page_end=page_count,
    )

    if not raw_chunks:
        raise DocumentProcessException("文档分块失败，未生成任何片段")

    # ===== 第3步：去重 =====
    seen_hashes = set()
    unique_chunks = []

    for chunk in raw_chunks:
        content_hash = hashlib.sha256(
            chunk["content"].encode("utf-8")
        ).hexdigest()
        if content_hash not in seen_hashes:
            seen_hashes.add(content_hash)
            unique_chunks.append(chunk)

    # ===== 第4步：生成 Embeddings =====
    chunk_texts = [c["content"] for c in unique_chunks]
    chunk_embeddings = embeddings.embed_documents(chunk_texts)

    if len(chunk_embeddings) != len(unique_chunks):
        raise DocumentProcessException(
            f"向量化数量不匹配: {len(chunk_embeddings)} != {len(unique_chunks)}"
        )

    # ===== 第5步：写入 ChromaDB =====
    chunk_ids = []
    chroma_ids = []
    chroma_metadatas = []

    for i, chunk in enumerate(unique_chunks):
        cid = f"doc_{document_id}_chunk_{i}"
        chunk_ids.append(cid)
        chroma_ids.append(cid)
        chroma_metadatas.append({
            "document_id": document_id,
            "document_name": title,
            "chapter": chunk.get("chapter", ""),
            "page": chunk.get("page", 1),
            "security_level": security_level,
            "category_id": category_id or 0,
            "token_count": chunk.get("token_count", 0),
            "chunk_index": i,
        })

    vector_store.add_embeddings(
        ids=chroma_ids,
        embeddings=chunk_embeddings,
        documents=chunk_texts,
        metadatas=chroma_metadatas,
    )

    # ===== 第6步：写入 BM25 索引 =====
    bm25_chunks = [
        {
            "chunk_id": cid,
            "document_id": document_id,
            "document_name": title,
            "chapter": chunk.get("chapter", ""),
            "page": chunk.get("page", 1),
            "content": chunk["content"],
            "security_level": security_level,
            "token_count": chunk.get("token_count", 0),
        }
        for cid, chunk in zip(chunk_ids, unique_chunks)
    ]
    bm25_retriever.add_chunks(bm25_chunks)

    # ===== 第7步：持久化 BM25 索引 =====
    save_bm25_index()

    # 构建 chunk 记录（含内容，供数据库存储）
    chunk_records = [
        {
            "content": chunk["content"],
            "chapter": chunk.get("chapter", ""),
            "page": chunk.get("page", 1),
            "token_count": chunk.get("token_count", 0),
        }
        for chunk in unique_chunks
    ]

    return len(unique_chunks), chunk_ids, chunk_records


def remove_document_from_index(document_id: int) -> int:
    """
    从向量库和 BM25 索引中移除文档的所有 chunk

    返回: 移除的 chunk 数量
    """
    # 从 ChromaDB 移除
    vector_store = get_vector_store()
    vector_store.delete_by_filter({"document_id": document_id})

    # 从 BM25 移除
    bm25_retriever = get_bm25_retriever()
    removed = bm25_retriever.remove_by_document_id(document_id)

    # 持久化
    save_bm25_index()

    return removed


# ==================== 辅助函数 ====================


def _table_to_markdown(table) -> str:
    """将 python-docx 表格转为 Markdown 格式"""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)

    if not rows:
        return ""

    return _rows_to_markdown_table(rows)


def _rows_to_markdown_table(rows: List[List[str]]) -> str:
    """将二维列表转为 Markdown 表格"""
    if not rows:
        return ""

    max_cols = max(len(row) for row in rows)
    padded_rows = [row + [""] * (max_cols - len(row)) for row in rows]

    lines = []
    # 表头
    lines.append("| " + " | ".join(padded_rows[0]) + " |")
    # 分隔线
    lines.append("| " + " | ".join(["---"] * max_cols) + " |")
    # 数据行
    for row in padded_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)
